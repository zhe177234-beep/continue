import io
import sys
import sqlite3
from pathlib import Path
import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'ai-service'))
from api import create_app
from engine import Engine, parse_document
from models import Ollama
from store import Store
from retrieval import reindex, retrieve


@pytest.fixture
def app(tmp_path):
    return create_app(tmp_path, Ollama(model='', embedding_model=''))


def account(app, name='student_one'):
    client = TestClient(app)
    credentials = {'username': name, 'password': 'safe-password-123'}
    assert client.post('/api/auth/register', json=credentials).status_code == 201
    result = client.post('/api/auth/login', json=credentials)
    assert result.status_code == 200
    assert 'HttpOnly' in result.headers['set-cookie']
    return client


def seeded(app):
    client = account(app)
    base = client.post('/api/bases', json={'name': '机器学习'}).json()['id']
    document = client.post(f'/api/bases/{base}/documents', files={'file': ('ml.md', (ROOT/'datasets/machine-learning.md').read_bytes())})
    assert document.status_code == 201, document.text
    return client, base, document.json()


def test_login_logout_password_and_session_storage(app):
    client = account(app)
    assert client.get('/api/auth/me').json()['username'] == 'student_one'
    with app.state.store.db() as db:
        row = db.execute('SELECT * FROM users').fetchone()
        assert row['password'] != 'safe-password-123'
        assert len(row['salt']) == 32
        token = db.execute('SELECT token FROM sessions').fetchone()[0]
        assert token != client.cookies['session']
    assert client.post('/api/auth/logout', json={}).status_code == 200
    assert client.get('/api/auth/me').status_code == 401
    assert client.post('/api/auth/login', json={'username':'student_one','password':'wrong-password'}).status_code == 401


def test_cross_account_knowledge_isolation(app):
    first, base, doc = seeded(app)
    second = account(app, 'student_two')
    assert second.get('/api/bases').json() == []
    for route in ['documents','graph','history','progress']:
        assert second.get(f'/api/bases/{base}/{route}').status_code == 404
    assert second.post(f'/api/bases/{base}/ask', json={'question':'梯度下降'}).status_code == 404
    assert second.delete(f'/api/bases/{base}/documents/{doc["id"]}').status_code == 404
    assert len(first.get(f'/api/bases/{base}/documents').json()) == 1


def test_ask_citations_history_and_refusal(app):
    client, base, _ = seeded(app)
    for mode in ['auto','bm25','cosine','hybrid','graph']:
        response = client.post(f'/api/bases/{base}/ask', json={'question':'学习率和梯度下降的关系','mode':mode})
        assert response.status_code == 200, response.text
        data = response.json()
        assert data['answer_kind'] == 'extractive'
        assert data['sources'][0]['name'] == 'ml.md'
        assert data['sources'][0]['page'] == 1
        assert '学习率' in data['answer']
    result = client.post(f'/api/bases/{base}/ask', json={'question':'火星移民交通费用'}).json()
    assert result['sources'] == []
    assert len(client.get(f'/api/bases/{base}/history').json()) == 6


def test_quiz_scoring_mastery_and_delete_cleanup(app):
    client, base, doc = seeded(app)
    quizzes = client.post(f'/api/bases/{base}/quizzes', json={'count':5}).json()
    assert len(quizzes) == 5
    assert all('answer' not in q for q in quizzes)
    q = quizzes[0]
    with app.state.store.db() as db:
        correct = db.execute('SELECT answer FROM quizzes WHERE id=?',(q['id'],)).fetchone()[0]
    wrong = (correct+1) % len(q['options'])
    result = client.post(f'/api/bases/{base}/quizzes/{q["id"]}/submit',json={'selected':wrong})
    assert result.status_code == 200
    assert result.json()['correct'] is False
    assert result.json()['answer'] == correct
    assert client.post(f'/api/bases/{base}/quizzes/{q["id"]}/submit',json={'selected':correct}).status_code == 400
    p = client.get(f'/api/bases/{base}/progress').json()
    assert len(p['mistakes']) == 1
    assert sum(x['attempts'] for x in p['mastery']) == 1
    assert client.delete(f'/api/bases/{base}/documents/{doc["id"]}').status_code == 200
    assert client.get(f'/api/bases/{base}/graph').json() == []
    assert client.get(f'/api/bases/{base}/progress').json()['mistakes'] == []


def test_duplicate_persistence(app, tmp_path):
    client, base, doc = seeded(app)
    result = client.post(f'/api/bases/{base}/documents',files={'file':('other.md',(ROOT/'datasets/machine-learning.md').read_bytes())}).json()
    assert result['duplicate'] is True and result['id'] == doc['id']
    new = Store(tmp_path)
    user = new.user(client.cookies['session'])
    assert new.owns(user['id'],base)
    assert len(Engine(tmp_path/'bases'/(base+'.db')).documents()) == 1


def test_validation_and_csrf(app):
    client, base, _ = seeded(app)
    assert client.post('/api/bases', json={'name':' '}).status_code == 422
    assert client.post('/api/bases',json={'name':'foreign'},headers={'Origin':'https://evil.example'}).status_code == 403
    for data in [{'question':''},{'question':'x','k':True},{'question':'x','mode':'bad'}]:
        assert client.post(f'/api/bases/{base}/ask',json=data).status_code == 422
    assert client.post(f'/api/bases/{base}/ask',json={'question':'x','mode':'semantic'}).status_code == 400
    assert client.post(f'/api/bases/{base}/documents',files={'file':('bad.txt',b'\xff')}).status_code == 400
    assert client.post(f'/api/bases/{base}/documents',files={'file':('large.txt',b'x'*(3*1024*1024+1))}).status_code in {400,413}


def test_office_parsers_and_corrupt_pdf():
    from docx import Document
    from pptx import Presentation
    document=Document();document.add_paragraph('Gradient descent optimizes loss.');buffer=io.BytesIO();document.save(buffer)
    assert 'Gradient' in parse_document('notes.docx',buffer.getvalue())[0][1]
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text='Training and test sets';buffer=io.BytesIO();presentation.save(buffer)
    pages=parse_document('slides.pptx',buffer.getvalue())
    assert pages[0][0] == 1 and 'Training' in pages[0][1]
    with pytest.raises(ValueError): parse_document('bad.pdf',b'not a PDF')
    with pytest.raises(ValueError): parse_document('bad.docx',b'not a ZIP')


class FakeModel:
    model='fake-chat';embedding_model='fake-embed'
    def embed(self,texts):
        return [[1.,0.] if '梯度' in t else [0.,1.] for t in texts]
    def generate(self,question,sources): return '梯度下降需要选择合适的学习率。[1]'


def test_semantic_index_and_generation(tmp_path):
    model=FakeModel();index=Engine(tmp_path/'index.db')
    index.ingest('a.txt','梯度下降需要学习率'.encode());index.ingest('b.txt','正则化减少过拟合'.encode())
    assert reindex(index,model)['indexed']==2
    result=retrieve(index,model,'梯度下降','semantic',3)
    assert result['sources'][0]['name']=='a.txt'
    app=create_app(tmp_path/'app',model);client,base,_=seeded(app)
    assert client.post(f'/api/bases/{base}/index',json={}).status_code==200
    assert client.post(f'/api/bases/{base}/ask',json={'question':'梯度下降'}).json()['answer_kind']=='generated'


def test_model_failure_falls_back(tmp_path):
    class Broken(FakeModel):
        def embed(self,texts): raise RuntimeError('unavailable')
        def generate(self,question,sources): raise RuntimeError('unavailable')
    app=create_app(tmp_path,Broken());client,base,_=seeded(app)
    result=client.post(f'/api/bases/{base}/ask',json={'question':'梯度下降'}).json()
    assert result['answer_kind']=='extractive'
    assert '降级' in result['warning']
    assert client.post(f'/api/bases/{base}/index',json={}).status_code==503


def test_model_citation_and_vector_contract(monkeypatch):
    model=Ollama(model='test',embedding_model='test')
    monkeypatch.setattr(model,'post',lambda *a: {'message':{'content':'虚构引用 [9]'}})
    with pytest.raises(ValueError): model.generate('x',[{'text':'evidence'}])
    monkeypatch.setattr(model,'post',lambda *a: {'embeddings':[[float('nan')]]})
    with pytest.raises(ValueError): model.embed(['x'])


def test_auth_rate_limit(app):
    client=TestClient(app)
    for _ in range(20):
        assert client.post('/api/auth/login',json={'username':'unknown','password':'not-a-valid-password'}).status_code==401
    assert client.post('/api/auth/login',json={'username':'unknown','password':'not-a-valid-password'}).status_code==429


def test_pdf_actual_pages(tmp_path):
    from reportlab.pdfgen import canvas
    buffer=io.BytesIO();pdf=canvas.Canvas(buffer)
    pdf.drawString(50,700,'Gradient descent depends on learning rate.');pdf.showPage()
    pdf.drawString(50,700,'Regularization reduces overfitting.');pdf.save()
    index=Engine(tmp_path/'pdf.db');index.ingest('lecture.pdf',buffer.getvalue())
    assert index.search('regularization')['sources'][0]['page']==2
    assert index.search('learning rate')['sources'][0]['page']==1


def test_real_image_ocr(monkeypatch):
    import shutil
    if not shutil.which('tesseract'): pytest.skip('Tesseract unavailable in this environment')
    from PIL import Image, ImageDraw, ImageFont
    monkeypatch.setenv('ENABLE_OCR','true');monkeypatch.setenv('OCR_LANG','eng')
    image=Image.new('RGB',(700,120),'white')
    try: font=ImageFont.truetype('DejaVuSans.ttf',36)
    except OSError: font=ImageFont.load_default(size=36)
    ImageDraw.Draw(image).text((20,35),'Gradient descent',fill='black',font=font)
    buffer=io.BytesIO();image.save(buffer,format='PNG')
    assert 'gradient' in parse_document('photo.png',buffer.getvalue())[0][1].lower()


def test_concurrent_duplicate_ingestion(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    index=Engine(tmp_path/'atomic.db')
    with ThreadPoolExecutor(max_workers=4) as pool:
        results=list(pool.map(lambda _:index.ingest('same.txt','梯度下降学习率'.encode()),range(4)))
    assert sum(not r['duplicate'] for r in results)==1
    assert len(index.documents())==1
