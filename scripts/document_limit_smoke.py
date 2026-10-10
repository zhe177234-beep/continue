"""Exercise real 1000-page PDF/PPTX uploads and reject page 1001.

Requires requirements-dev.txt (ReportLab generates the PDF fixtures).
Only an isolated verification account is used; its documents are removed.
"""
import argparse
import io
import time

import httpx
from pptx import Presentation
from reportlab.pdfgen import canvas


def run(url):
    with httpx.Client(base_url=url.rstrip('/'),timeout=210,trust_env=False) as client:
        limits=client.get('/api/health').json()['limits']
        assert limits=={'upload_bytes':100*1024*1024,'document_pages':1000,'ocr_pages':20},limits
        credentials={'username':f'pages_{int(time.time()*1000)}','password':'document-pages-smoke-123'}
        client.post('/api/auth/register',json=credentials).raise_for_status()
        client.post('/api/auth/login',json=credentials).raise_for_status()
        response=client.post('/api/bases',json={'name':'Document page limit verification'});response.raise_for_status()
        prefix='/api/bases/'+response.json()['id'];documents=[]
        try:
            for suffix in ('pdf','pptx'):
                for count in (1000,1001):
                    buffer=io.BytesIO()
                    if suffix=='pdf':
                        pdf=canvas.Canvas(buffer,pageCompression=1)
                        for page in range(1,count+1):
                            pdf.drawString(40,700,f'BoundaryToken{page} learning notes.');pdf.showPage()
                        pdf.save()
                    else:
                        presentation=Presentation()
                        for page in range(1,count+1):
                            slide=presentation.slides.add_slide(presentation.slide_layouts[5])
                            slide.shapes.title.text=f'BoundaryToken{page} learning notes.'
                        presentation.save(buffer)
                    response=client.post(prefix+'/documents',files={'file':(f'pages-{count}.{suffix}',buffer.getvalue())})
                    if count==1001:
                        assert response.status_code==400 and '1000' in response.json()['detail'],response.text
                        print(f'PASS {suffix.upper()} page 1001 rejected through full stack',flush=True)
                    else:
                        response.raise_for_status();document_id=response.json()['id'];documents.append(document_id)
                        assert response.json()['chunks']==1000,response.text
                        answer=client.post(prefix+'/ask',json={'question':'BoundaryToken1000','mode':'bm25','generate':False})
                        answer.raise_for_status()
                        assert any(source['document_id']==document_id and source['page']==1000 for source in answer.json()['sources']),answer.text
                        print(f'PASS actual 1000-page {suffix.upper()} uploaded and final-page citation preserved',flush=True)
        finally:
            for document_id in documents:
                client.delete(prefix+'/documents/'+document_id).raise_for_status()
            client.post('/api/auth/logout',json={}).raise_for_status()


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--url',default='http://localhost:8080')
    run(parser.parse_args().url)
