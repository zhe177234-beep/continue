"""Offline retrieval baseline. TF cosine is lexical, not semantic embedding."""
import collections
import hashlib
import io
import math
import re
import sqlite3
import os
import zipfile
import subprocess
import tempfile
from pathlib import Path

MAX_FILE = 10 * 1024 * 1024


def tokens(text):
    # Chinese bigrams preserve more meaning than isolated characters.
    result = re.findall(r"[a-z0-9_]+", text.lower())
    for run in re.findall(r"[\u4e00-\u9fff]+", text):
        result.extend(run[i:i + 2] for i in range(len(run) - 1))
        if len(run) == 1:
            result.append(run)
    return result


def parse_document(name, data):
    if not data or len(data) > MAX_FILE:
        raise ValueError("文件为空或超过 10 MiB")
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md"}:
        try:
            pages = [(1, data.decode("utf-8-sig"))]
        except UnicodeDecodeError:
            raise ValueError("文本必须使用 UTF-8 编码") from None
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("PDF 支持需要先安装 requirements.txt") from None
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("不支持加密 PDF")
            if len(reader.pages) > 100:
                raise ValueError("PDF 不能超过 100 页")
            pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
            blank = [i for i, (_,text) in enumerate(pages) if not text.strip()]
            if blank and os.getenv('ENABLE_OCR', 'false').lower() == 'true':
                if len(blank)>20: raise ValueError('扫描 PDF 每次最多 OCR 20 页，请拆分文件')
                from PIL import Image
                import pytesseract
                with tempfile.TemporaryDirectory(prefix='zhixue-pdf-') as folder:
                    pdf_path = Path(folder)/'source.pdf'
                    pdf_path.write_bytes(data)
                    for i in blank:
                        output = Path(folder)/'page'
                        subprocess.run(['pdftoppm','-f',str(i+1),'-l',str(i+1),'-singlefile','-scale-to','1800','-png',str(pdf_path),str(output)],check=True,timeout=20,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
                        with Image.open(str(output)+'.png') as image:
                            text = pytesseract.image_to_string(image,lang=os.getenv('OCR_LANG','chi_sim+eng'),timeout=20)
                        pages[i]=(i+1,text)

        except ValueError:
            raise
        except Exception:
            raise ValueError("PDF 解析或扫描页 OCR 失败，请检查文件、Poppler 和 OCR 语言包") from None
    elif suffix in {'.docx', '.pptx'}:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 5000 or sum(x.file_size for x in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError('Office 文件解压后过大')
            if suffix == '.docx':
                from docx import Document
                document = Document(io.BytesIO(data))
                text = '\n'.join(p.text for p in document.paragraphs)
                text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for table in document.tables for row in table.rows)
                pages = [(1, text)]
            else:
                from pptx import Presentation
                presentation = Presentation(io.BytesIO(data))
                if len(presentation.slides) > 100: raise ValueError('PPT 不能超过 100 页')
                pages = [(i + 1, '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)) for i, slide in enumerate(presentation.slides)]
        except ValueError:
            raise
        except Exception:
            raise ValueError('Office 文档解析失败，仅支持 DOCX/PPTX') from None
    elif suffix in {'.png', '.jpg', '.jpeg'}:
        if os.getenv('ENABLE_OCR', 'false').lower() != 'true':
            raise ValueError('图片 OCR 未启用；设置 ENABLE_OCR=true 并安装 Tesseract')
        try:
            from PIL import Image
            import pytesseract
            image = Image.open(io.BytesIO(data))
            if image.width * image.height > 16_000_000: raise ValueError('图片像素过大')
            text = pytesseract.image_to_string(image, lang=os.getenv('OCR_LANG', 'chi_sim+eng'), timeout=20)
            pages = [(1, text)]
        except ValueError:
            raise
        except Exception:
            raise ValueError('OCR 失败，请检查图片、Tesseract 和语言包') from None
    else:
        raise ValueError("支持 TXT、Markdown、PDF、DOCX、PPTX；图片需启用 OCR")
    if sum(len(t) for _, t in pages) > 1_000_000:
        raise ValueError("解析后的文本过大")
    if not any(t.strip() for _, t in pages):
        raise ValueError("未提取到文本；扫描 PDF 需启用 OCR 并安装 Poppler/Tesseract")
    return pages


class Engine:
    def __init__(self, path):
        self.path = str(path)
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS documents(
                  id TEXT PRIMARY KEY, name TEXT NOT NULL, digest TEXT UNIQUE NOT NULL,
                  created TEXT DEFAULT CURRENT_TIMESTAMP);
                CREATE TABLE IF NOT EXISTS chunks(
                  id TEXT PRIMARY KEY, document_id TEXT REFERENCES documents(id) ON DELETE CASCADE,
                  page INTEGER NOT NULL, position INTEGER NOT NULL, text TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS relations(
                  chunk_id TEXT REFERENCES chunks(id) ON DELETE CASCADE,
                  subject TEXT, predicate TEXT, object TEXT);
                CREATE TABLE IF NOT EXISTS embeddings(
                  chunk_id TEXT PRIMARY KEY REFERENCES chunks(id) ON DELETE CASCADE,
                  model TEXT NOT NULL, vector TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS graph_meta(key TEXT PRIMARY KEY,value TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS graph_entities(
                  id TEXT PRIMARY KEY,name TEXT NOT NULL,aliases TEXT NOT NULL,
                  chunk_ids TEXT NOT NULL,description TEXT NOT NULL,vector TEXT);
                CREATE TABLE IF NOT EXISTS graph_communities(
                  id TEXT PRIMARY KEY,level INTEGER NOT NULL,parent_id TEXT,
                  entity_ids TEXT NOT NULL,chunk_ids TEXT NOT NULL,report TEXT NOT NULL);
            """)

    def connect(self):
        db = sqlite3.connect(self.path, timeout=15)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        return db

    def ingest(self, name, data):
        name = name.replace("\\", "/").split("/")[-1]
        if not name or len(name) > 160:
            raise ValueError("文件名不能为空或超过 160 字符")
        pages = parse_document(name, data)
        digest = hashlib.sha256(data).hexdigest()
        doc_id = digest[:24]
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            old = db.execute("SELECT id FROM documents WHERE digest=?", (digest,)).fetchone()
            if old:
                return {"id": old["id"], "duplicate": True}
            chunks = []
            for page, text in pages:
                text = text.replace("\x00", "").strip()
                # 100-character overlap; page boundaries are never crossed.
                for position, offset in enumerate(range(0, len(text), 500)):
                    part = text[offset:offset + 600].strip()
                    if part:
                        chunks.append((f"{doc_id}:{page}:{position}", doc_id, page, position, part))
            count = db.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
            if count + len(chunks) > 2000:
                raise ValueError("演示知识库最多容纳 2000 个分块，请删除旧资料后重试")
            db.execute("INSERT INTO documents(id,name,digest) VALUES(?,?,?)", (doc_id, name, digest))
            db.executemany("INSERT INTO chunks VALUES(?,?,?,?,?)", chunks)
            self.invalidate_graph(db)
            # Explicit triples only; never invent automatic semantic relations.
            for chunk_id, _, _, _, text in chunks:
                for line in text.splitlines():
                    match = re.fullmatch(r"\s*关系[：:]\s*([^|\n]{1,60})\|([^|\n]{1,30})\|([^|\n]{1,60})\s*", line)
                    if match:
                        db.execute("INSERT INTO relations VALUES(?,?,?,?)", (chunk_id, *(s.strip() for s in match.groups())))
        return {"id": doc_id, "duplicate": False, "chunks": len(chunks)}

    def documents(self):
        with self.connect() as db:
            return [dict(r) for r in db.execute("SELECT d.id,d.name,d.created,COUNT(c.id) AS chunks FROM documents d LEFT JOIN chunks c ON c.document_id=d.id GROUP BY d.id ORDER BY d.created DESC,d.id")]

    def delete(self, doc_id):
        with self.connect() as db:
            deleted = db.execute("DELETE FROM documents WHERE id=?", (doc_id,)).rowcount > 0
            if deleted: self.invalidate_graph(db)
            return deleted

    @staticmethod
    def invalidate_graph(db):
        # Reports contain quotations, so deletion must remove derived copies too.
        db.execute('DELETE FROM graph_meta')
        db.execute('DELETE FROM graph_entities')
        db.execute('DELETE FROM graph_communities')

    def graph(self):
        with self.connect() as db:
            return [dict(r) for r in db.execute("SELECT DISTINCT subject,predicate,object,chunk_id FROM relations ORDER BY subject,predicate,object")]

    def search(self, question, mode="hybrid", k=5):
        if not isinstance(question, str) or not question.strip() or len(question) > 1000:
            raise ValueError("问题须为 1 至 1000 字符")
        if not isinstance(mode, str) or mode not in {"bm25", "cosine", "hybrid", "graph", "auto"}:
            raise ValueError("无效检索模式")
        if type(k) is not int or not 1 <= k <= 10:
            raise ValueError("k 必须为 1 至 10 的整数")
        if mode == "auto":
            mode = "graph" if any(w in question for w in ["关系", "关联", "依赖", "前置"]) else "hybrid"
        with self.connect() as db:
            rows = [dict(r) for r in db.execute("SELECT c.*,d.name FROM chunks c JOIN documents d ON d.id=c.document_id ORDER BY c.id")]
            relations = [dict(r) for r in db.execute("SELECT * FROM relations")]
        if not rows:
            return {"mode": mode, "sources": []}
        query = collections.Counter(tokens(question))
        corpus = [collections.Counter(tokens(r["text"])) for r in rows]
        df = collections.Counter(t for c in corpus for t in c)
        avglen = sum(sum(c.values()) for c in corpus) / len(corpus) or 1
        qnorm = math.sqrt(sum(v * v for v in query.values())) or 1
        bm, cosine = [], []
        for c in corpus:
            length = sum(c.values())
            score = sum(math.log(1 + (len(rows) - df[t] + .5) / (df[t] + .5)) * c[t] * 2.5 / (c[t] + 1.5 * (.25 + .75 * length / avglen)) for t in query if c[t])
            bm.append(score)
            cosine.append(sum(v * c[t] for t, v in query.items()) / (qnorm * (math.sqrt(sum(v*v for v in c.values())) or 1)))
        def ranks(values):
            return {idx: rank + 1 for rank, idx in enumerate(sorted(range(len(values)), key=lambda i: (-values[i], rows[i]["id"]))) if values[idx] > 0}
        br, cr = ranks(bm), ranks(cosine)
        scores = [sum(1 / (60 + r[i]) for r in [br, cr] if i in r) for i in range(len(rows))]
        if mode == "bm25":
            scores = bm
        elif mode == "cosine":
            scores = cosine
        elif mode == "graph":
            seeds = {r[key] for r in relations for key in ["subject", "object"] if r[key] in question}
            evidence_ids = set()
            frontier = seeds
            visited = set(seeds)
            # Bounded two-hop expansion, retaining source chunks for every edge.
            for _ in range(2):
                edges = [r for r in relations if r["subject"] in frontier or r["object"] in frontier]
                evidence_ids.update(r["chunk_id"] for r in edges)
                adjacent = {r[key] for r in edges for key in ("subject", "object")}
                frontier = adjacent - visited
                visited.update(adjacent)
            scores = [s + (.1 if row["id"] in evidence_ids else 0) for row, s in zip(rows, scores)]
        ordered = sorted(range(len(rows)), key=lambda i: (-scores[i], rows[i]["id"]))
        sources = []
        for i in ordered:
            # Weak lexical overlap alone is insufficient evidence.
            if scores[i] <= 0 or (mode != "graph" and cosine[i] < .05):
                continue
            row = rows[i]
            sources.append({"chunk_id": row["id"], "document_id": row["document_id"], "name": row["name"], "page": row["page"], "text": row["text"], "score": round(scores[i], 6)})
            if len(sources) == k:
                break
        return {"mode": mode, "sources": sources}

    def answer(self, question, mode="auto", k=5):
        result = self.search(question, mode, k)
        result["answer_kind"] = "extractive"
        result["answer"] = "\n\n".join(f"[{i}] {s['text']}" for i, s in enumerate(result["sources"], 1)) if result["sources"] else "资料中没有找到足够相关的证据，请补充资料或更具体地提问。"
        return result
